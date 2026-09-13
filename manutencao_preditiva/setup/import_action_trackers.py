# Copyright (c) 2026, Dércio Bobo and contributors
# For license information, please see license.txt
"""One-off import of the legacy "Action Tracker" Excel files (vibration
analysis and thermography) into Campanha De Inspecao / Achado De Inspecao.

Not wired into hooks.py/patches.txt on purpose - this is historical data
for a specific rollout, not a schema migration every install should run.

The Excel files themselves are never committed (see .gitignore) - they were
extracted locally (wherever openpyxl is available) into a single plain,
frappe-independent JSON file that IS committed with the app:
setup/data/achados_import.json. That means the server never needs openpyxl
or the original spreadsheets at all - once the app is deployed there via
the usual `bench get-app`/git pull, just run:

	bench --site <site> execute manutencao_preditiva.setup.import_action_trackers.load_from_json

To regenerate setup/data/achados_import.json after a source file changes,
run this locally (needs openpyxl, not frappe):

	python -c "
	from manutencao_preditiva.setup.import_action_trackers import extract_to_json
	extract_to_json(
	    vibration_file='FR.TEC.014...xlsx',
	    thermography_file='FR.TEC.016...xlsx',
	)"

(`run()` below does both steps in one call, for the case where openpyxl
*is* installed on the bench - kept for convenience/testing.)
"""

import json
import os

VIBRATION_SEVERITY = {
	1: "Crítico",
	2: "Alarme",
	3: "Boa Condição",
	4: "Aceitável",
	5: "Não Recolhido",
}

# Cell fill colours actually used as severity highlights in the thermography
# tracker (the legend's own default/theme banding colour is deliberately
# NOT in here - only manually-applied highlights mean anything).
THERMO_CRITICAL_RGB = {"FFFF0000"}
THERMO_ALARM_RGB = {"FFFFC000", "FFED7D31", "FFFFA500"}
THERMO_LOW_RGB = {"FFFFFF00"}

ACTION_STATUS_MAP = {
	"": "Pendente",
	"OPEN": "Pendente",
	"PENDING": "Pendente",
	"IN PROGRESS": "Em Curso",
	"ONGOING": "Em Curso",
	"CLOSED": "Concluído",
	"DONE": "Concluído",
	"COMPLETE": "Concluído",
	"COMPLETED": "Concluído",
	"N/A": "Não Aplicável",
	"NA": "Não Aplicável",
}


# ---------------------------------------------------------------------------
# Step 1: extraction (needs openpyxl, does NOT need frappe) - run locally.
# ---------------------------------------------------------------------------


DEFAULT_JSON_PATH = os.path.join(os.path.dirname(__file__), "data", "achados_import.json")


def extract_to_json(vibration_file=None, thermography_file=None, output_path=None,
	vibration_cliente="Kenmare", thermography_cliente="CLN"):
	"""Read the Excel file(s) and write one flat, frappe-independent JSON
	file describing the Campanhas and Achados to create."""
	output_path = output_path or DEFAULT_JSON_PATH
	campanhas = []
	achados = []

	if vibration_file:
		campanha, rows = extract_vibration(vibration_file, vibration_cliente)
		campanhas.append(campanha)
		achados.extend(rows)

	if thermography_file:
		campanha, rows = extract_thermography(thermography_file, thermography_cliente)
		campanhas.append(campanha)
		achados.extend(rows)

	payload = {"campanhas": campanhas, "achados": achados}
	with open(output_path, "w", encoding="utf-8") as f:
		json.dump(payload, f, ensure_ascii=False, indent=1, default=_json_default)

	print(f"Escrito {output_path}: {len(campanhas)} campanha(s), {len(achados)} achado(s)")
	return payload


def extract_vibration(file_path, cliente_nome):
	import openpyxl

	wb = openpyxl.load_workbook(file_path, data_only=True)

	campanha_key = "vibracao"
	campanha = {
		"key": campanha_key,
		"cliente": cliente_nome,
		"tecnica": "Vibração",
		"data_da_inspecao": None,
		"referencia_do_documento": "FR.TEC.014",
	}
	achados = []

	for sheet_name in wb.sheetnames:
		if sheet_name.strip().upper() == "DROPDOWN MENU":
			continue

		ws = wb[sheet_name]
		header_row = _find_header_row(ws, expected_label="Plant")
		if not header_row:
			continue

		if campanha["data_da_inspecao"] is None:
			campanha["data_da_inspecao"] = ws.cell(row=header_row, column=6).value

		for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
			equipamento = _cellval(row, 3)
			if not equipamento:
				continue

			severidade_raw = _cellval(row, 6)
			try:
				severidade = VIBRATION_SEVERITY.get(int(severidade_raw), "Não Recolhido")
			except (TypeError, ValueError):
				severidade = "Não Recolhido"

			achados.append(
				{
					"campanha_key": campanha_key,
					"area_planta": _cellval(row, 4) or sheet_name,
					"item": _cellval(row, 1),
					"equipamento_referencia": equipamento,
					"equipamento_descricao": _cellval(row, 5),
					"componente": _cellval(row, 8),
					"severidade": severidade,
					"descricao_do_defeito": _cellval(row, 7) or "-",
					"acao_recomendada": _cellval(row, 9),
					"resposta_do_cliente": _cellval(row, 10),
				}
			)

	return campanha, achados


def extract_thermography(file_path, cliente_nome):
	import openpyxl

	wb = openpyxl.load_workbook(file_path, data_only=True)

	campanha_key = "termografia"
	campanha = {
		"key": campanha_key,
		"cliente": cliente_nome,
		"tecnica": "Termografia",
		"data_da_inspecao": None,
		"referencia_do_documento": "FR.TEC.016",
	}
	achados = []

	for sheet_name in wb.sheetnames:
		ws = wb[sheet_name]
		header_row = _find_header_row(ws, expected_label="Item")
		if not header_row:
			continue

		for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
			item = _cellval(row, 2)
			if item is None:
				continue

			equipamento = _cellval(row, 4)
			if not equipamento:
				continue

			temp_max = _to_float(_cellval(row, 10))
			temp_actual = _to_float(_cellval(row, 11))
			temp_amb = _to_float(_cellval(row, 12))
			temp_diff = _to_float(_cellval(row, 13))

			severidade = _classify_thermo_severity(row[3])  # Equipment Reference cell, column D

			descricao = (
				f"Ponto quente detectado por termografia: {temp_actual}°C "
				f"(máx. operação {temp_max}°C, ambiente {temp_amb}°C, diferença {temp_diff}°C)."
				if temp_actual is not None
				else "Ponto quente detectado por termografia."
			)

			achados.append(
				{
					"campanha_key": campanha_key,
					"area_planta": f"{sheet_name} - {_cellval(row, 3)}" if _cellval(row, 3) else sheet_name,
					"item": item,
					"equipamento_referencia": equipamento,
					"equipamento_descricao": _cellval(row, 5),
					"componente": _cellval(row, 7),
					"numero_da_imagem": _cstr(_cellval(row, 8)),
					"ordem_de_servico": _cstr(_cellval(row, 9)),
					"severidade": severidade,
					"descricao_do_defeito": descricao,
					"acao_recomendada": _cellval(row, 14),
					"plano_de_monitorizacao": _cellval(row, 15),
					"temp_max_operacao": temp_max,
					"temp_actual": temp_actual,
					"temp_ambiente": temp_amb,
					"diferenca_sobre_max": temp_diff,
					"resposta_do_cliente": _cellval(row, 16),
					"responsavel": _cellval(row, 17),
					"prazo": _cellval(row, 18),
					"estado_da_accao": ACTION_STATUS_MAP.get(
						_cstr(_cellval(row, 19)).strip().upper(), "Pendente"
					),
					"data_de_conclusao": _cellval(row, 20),
				}
			)

	return campanha, achados


# ---------------------------------------------------------------------------
# Step 2: loading (needs frappe, does NOT need openpyxl) - run on the site.
# ---------------------------------------------------------------------------


def load_from_json(json_path=None):
	import frappe

	json_path = json_path or DEFAULT_JSON_PATH
	with open(json_path, encoding="utf-8") as f:
		payload = json.load(f)

	created = {"campanhas": 0, "achados": 0}
	campanha_name_by_key = {}

	for campanha in payload["campanhas"]:
		_ensure_customer(campanha["cliente"])
		name = _ensure_campanha(
			cliente=campanha["cliente"],
			tecnica=campanha["tecnica"],
			data_da_inspecao=campanha["data_da_inspecao"],
			referencia_do_documento=campanha["referencia_do_documento"],
			created=created,
		)
		campanha_name_by_key[campanha["key"]] = name

	for achado in payload["achados"]:
		fields = {**achado, "campanha": campanha_name_by_key[achado["campanha_key"]]}
		fields.pop("campanha_key")
		if _create_achado_if_new(fields):
			created["achados"] += 1

	frappe.db.commit()
	print(f"Campanhas criadas: {created['campanhas']}  Achados criados: {created['achados']}")
	return created


def run(vibration_file=None, thermography_file=None, vibration_cliente="Kenmare", thermography_cliente="CLN"):
	"""Convenience one-shot path for when openpyxl IS available on the
	bench: extract straight to a temp JSON, then load it immediately."""
	import tempfile

	with tempfile.NamedTemporaryFile(suffix=".json", delete=False) as tmp:
		tmp_path = tmp.name

	extract_to_json(
		vibration_file=vibration_file,
		thermography_file=thermography_file,
		output_path=tmp_path,
		vibration_cliente=vibration_cliente,
		thermography_cliente=thermography_cliente,
	)
	return load_from_json(tmp_path)


def _ensure_customer(nome):
	import frappe

	if frappe.db.exists("Customer", nome):
		return nome

	# Selling Settings' defaults (or the hardcoded "All ..." root names) can
	# point at a *group* node, which ERPNext's Customer.validate() rejects -
	# a Customer must sit on a non-group leaf. Fall back to any leaf record.
	customer_group = frappe.db.get_single_value("Selling Settings", "customer_group")
	if not customer_group or frappe.db.get_value("Customer Group", customer_group, "is_group"):
		customer_group = frappe.db.get_value("Customer Group", {"is_group": 0})

	territory = frappe.db.get_single_value("Selling Settings", "territory")
	if not territory or frappe.db.get_value("Territory", territory, "is_group"):
		territory = frappe.db.get_value("Territory", {"is_group": 0})

	doc = frappe.get_doc(
		{
			"doctype": "Customer",
			"customer_name": nome,
			"customer_group": customer_group,
			"territory": territory,
		}
	)
	doc.insert(ignore_permissions=True, ignore_mandatory=True)
	return doc.name


def _ensure_campanha(cliente, tecnica, data_da_inspecao, referencia_do_documento, created):
	import frappe

	existing = frappe.db.get_value(
		"Campanha De Inspecao",
		{"cliente": cliente, "tecnica": tecnica, "referencia_do_documento": referencia_do_documento},
	)
	if existing:
		return existing

	doc = frappe.get_doc(
		{
			"doctype": "Campanha De Inspecao",
			"cliente": cliente,
			"tecnica": tecnica,
			"data_da_inspecao": data_da_inspecao,
			"referencia_do_documento": referencia_do_documento,
		}
	)
	doc.insert(ignore_permissions=True)
	created["campanhas"] += 1
	return doc.name


def _create_achado_if_new(fields):
	import frappe

	dedup_filters = {
		"campanha": fields["campanha"],
		"area_planta": fields["area_planta"],
		"equipamento_referencia": fields["equipamento_referencia"],
		"item": fields.get("item"),
	}
	if frappe.db.exists("Achado De Inspecao", dedup_filters):
		return False

	doc = frappe.get_doc({"doctype": "Achado De Inspecao", **fields})
	doc.insert(ignore_permissions=True)
	return True


# ---------------------------------------------------------------------------
# Shared helpers (no frappe, no openpyxl-specific types beyond the Cell
# passed into _classify_thermo_severity).
# ---------------------------------------------------------------------------


def _find_header_row(ws, expected_label, max_scan_rows=30, search_columns=(1, 2, 3)):
	for r in range(1, min(max_scan_rows, ws.max_row) + 1):
		for c in search_columns:
			if _cstr(ws.cell(row=r, column=c).value).strip() == expected_label:
				return r
	return None


def _cellval(row, col_1_indexed):
	value = row[col_1_indexed - 1].value
	if isinstance(value, str):
		value = value.strip() or None
	return value


def _cstr(value):
	return "" if value is None else str(value)


def _to_float(value):
	if value in (None, ""):
		return None
	return float(value)


def _classify_thermo_severity(cell):
	fg = cell.fill.fgColor
	if fg is None:
		return "Boa Condição"

	if fg.type == "rgb":
		rgb = fg.rgb
		if rgb in THERMO_CRITICAL_RGB:
			return "Crítico"
		if rgb in THERMO_ALARM_RGB:
			return "Alarme"
		if rgb in THERMO_LOW_RGB:
			return "Aceitável"

	return "Boa Condição"


def _json_default(value):
	if hasattr(value, "isoformat"):
		return value.isoformat()
	return str(value)
